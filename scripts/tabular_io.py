#!/usr/bin/env python3
"""Load CSV, JSON and optionally XLSX as tabular rows, without guessing field meanings."""
import csv
import io
import json
from pathlib import Path


def _header(headers):
    if not headers or any(not isinstance(x, str) or not x.strip() for x in headers):
        raise ValueError("열 이름이 없거나 비어 있습니다")
    if len(headers) != len(set(headers)):
        raise ValueError("중복된 열 이름이 있습니다")
    return headers


def read_table(path, max_rows=100000, sheet=None):
    """Return (rows, headers). Reject oversized inputs instead of silently truncating."""
    path = Path(path)
    if max_rows < 1:
        raise ValueError("max_rows는 1 이상이어야 합니다")
    suffix = path.suffix.lower()
    if suffix == '.csv':
        raw = path.read_bytes()
        text = None
        for encoding in ('utf-8-sig', 'cp949'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            raise ValueError("CSV를 UTF-8 또는 CP949로 해석할 수 없습니다")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = _header(reader.fieldnames)
        rows = []
        for row in reader:
            if None in row:
                raise ValueError("열 수보다 값이 많은 CSV 행이 있습니다")
            rows.append(row)
            if len(rows) > max_rows:
                raise ValueError(f"허용 행 수({max_rows})를 초과했습니다")
        return rows, headers
    if suffix == '.json':
        data = json.loads(path.read_text(encoding='utf-8-sig'))
        if isinstance(data, dict):
            data = [data]
        if not isinstance(data, list) or not all(isinstance(row, dict) for row in data):
            raise ValueError("JSON은 객체 또는 객체 배열이어야 합니다")
        if len(data) > max_rows:
            raise ValueError(f"허용 행 수({max_rows})를 초과했습니다")
        headers = _header(list(dict.fromkeys(k for row in data for k in row))) if data else []
        return data, headers
    if suffix == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX 읽기에는 openpyxl 설치가 필요합니다") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet and sheet not in wb.sheetnames:
                raise ValueError(f"시트를 찾을 수 없습니다: {sheet}")
            ws = wb[sheet] if sheet else wb.active
            values = ws.iter_rows(values_only=True)
            header = next(values, None)
            headers = _header([str(value).strip() if value is not None else '' for value in header]) if header else []
            rows = []
            for values_row in values:
                if all(value is None for value in values_row):
                    continue
                if len(values_row) > len(headers) and any(v is not None for v in values_row[len(headers):]):
                    raise ValueError("헤더 열보다 값이 많은 XLSX 행이 있습니다")
                rows.append(dict(zip(headers, values_row)))
                if len(rows) > max_rows:
                    raise ValueError(f"허용 행 수({max_rows})를 초과했습니다")
            return rows, headers
        finally:
            wb.close()
    raise ValueError("지원 형식: .csv, .json, .xlsx (XLSX는 openpyxl 필요)")
